from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from typing import Any, Dict
import io
import pdfplumber
from docx import Document
import openpyxl

from ..azure_llm_client import generate_flow_sheets_from_text
from ..schemas_ai import GenerateSheetsRequest, GenerateSheetsResponse

router = APIRouter(
    prefix="/api/ai",
    tags=["ai"],
)

# ---------------------------------------------------------
# 기존 텍스트 기반 AI (그대로 유지)
# ---------------------------------------------------------
@router.post("/generate-sheets", response_model=GenerateSheetsResponse)
async def generate_sheets(req: GenerateSheetsRequest):
    sheets = await generate_flow_sheets_from_text(
        company_id=req.company_id,
        task_id=req.task_id,
        plain_text=req.plain_text,
    )
    return {
        "company_id": req.company_id,
        "task_id": req.task_id,
        "sheets": sheets,
    }


# ---------------------------------------------------------
# 파일 → 텍스트 변환 함수들
# ---------------------------------------------------------
async def extract_text_from_file(upload: UploadFile) -> str:
    filename = upload.filename or ""
    content_type = upload.content_type or ""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    raw_bytes = await upload.read()
    if not raw_bytes:
        raise HTTPException(status_code=400, detail="빈 파일입니다.")

    # TXT
    if ext == "txt" or content_type.startswith("text/"):
        try:
            return raw_bytes.decode("utf-8", errors="ignore")
        except Exception:
            return raw_bytes.decode("cp949", errors="ignore")

    # PDF
    if ext == "pdf" or content_type == "application/pdf":
        try:
            with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
                pages = [page.extract_text() or "" for page in pdf.pages]
            return "\n\n".join(pages)
        except Exception:
            raise HTTPException(
                status_code=500,
                detail="PDF 파일에서 텍스트를 추출하는 중 오류가 발생했습니다."
            )

    # DOCX
    if ext == "docx" or content_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ):
        try:
            doc = Document(io.BytesIO(raw_bytes))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            raise HTTPException(
                status_code=500,
                detail="Word(DOCX) 파일에서 텍스트 추출 중 오류가 발생했습니다."
            )

    # XLSX
    if ext == "xlsx" or content_type in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(c) for c in row if c is not None]
                    if row_vals:
                        rows.append("\t".join(row_vals))
            return "\n".join(rows)
        except Exception:
            raise HTTPException(
                status_code=500,
                detail="Excel(XLSX) 파일에서 텍스트 추출 중 오류가 발생했습니다."
            )

# ---------------------------------------------------------
# 파일 기반 AI: PDF / DOCX / XLSX / TXT → 텍스트 추출 → LLM 호출
# ---------------------------------------------------------
@router.post("/generate-sheets-from-file", response_model=GenerateSheetsResponse)
async def generate_sheets_from_file(
    company_id: str = Form(...),
    task_id: str = Form(...),
    file: UploadFile = File(...),
):
    # 1. 파일에서 텍스트 추출
    plain_text = await extract_text_from_file(file)

    # 2. LLM 호출
    sheets = await generate_flow_sheets_from_text(
        company_id=company_id,
        task_id=task_id,
        plain_text=plain_text,
    )

    return {
        "company_id": company_id,
        "task_id": task_id,
        "sheets": sheets,
    }
